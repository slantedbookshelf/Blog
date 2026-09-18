from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("outputs/documents")
OUT_PATH = OUT_DIR / "project-management-system-prd.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 38, 51)
MUTED = RGBColor(92, 102, 115)
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "D7DBE2"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_para_format(paragraph, before=0, after=6, line=1.10, alignment=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def paragraph_border_bottom(paragraph, color="2E74B5", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def add_p(doc, text="", style=None, before=0, after=6, line=1.10, bold=False, color=INK, size=11, align=None):
    p = doc.add_paragraph(style=style)
    set_para_format(p, before, after, line, align)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullets(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_para_format(p, after=4, line=1.167)
        if p.runs:
            p.runs[0].text = item
            set_run_font(p.runs[0], size=11, color=INK)
        else:
            set_run_font(p.add_run(item), size=11, color=INK)


def add_numbered(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_para_format(p, after=4, line=1.167)
        set_run_font(p.add_run(item), size=11, color=INK)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], header_fill=LIGHT_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        set_cell_shading(hdr[idx], header_fill)
        set_cell_border(hdr[idx])
        hdr[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[idx].paragraphs[0]
        set_para_format(p, after=0, line=1.0)
        r = p.add_run(text)
        set_run_font(r, size=9.5, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            set_cell_border(cells[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cells[idx].paragraphs[0]
            set_para_format(p, after=0, line=1.05)
            r = p.add_run(text)
            set_run_font(r, size=9.2, color=INK)
    set_cell_margins(table)
    set_table_geometry(table, widths)
    add_p(doc, "", after=2)
    return table


def add_label_detail_table(doc, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT_FILL)
        for cell in cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        p0 = cells[0].paragraphs[0]
        set_para_format(p0, after=0)
        set_run_font(p0.add_run(label), size=9.5, bold=True, color=INK)
        p1 = cells[1].paragraphs[0]
        set_para_format(p1, after=0)
        set_run_font(p1.add_run(value), size=9.5, color=INK)
    set_cell_margins(table)
    set_table_geometry(table, [2700, 6660])
    add_p(doc, "", after=2)


def add_callout(doc, title: str, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_border(cell, color="C8D0DA")
    p = cell.paragraphs[0]
    set_para_format(p, after=2)
    set_run_font(p.add_run(title + "："), size=10.5, bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(text), size=10.5, color=INK)
    set_cell_margins(table, top=120, bottom=120, start=160, end=160)
    set_table_geometry(table, [9360])
    add_p(doc, "", after=3)


def add_masthead(doc):
    header = doc.sections[0].header
    p = header.paragraphs[0]
    p.text = ""
    set_para_format(p, after=0)
    set_run_font(p.add_run("项目管理系统 PRD"), size=9, color=MUTED)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = doc.sections[0].footer
    pf = footer.paragraphs[0]
    pf.text = ""
    set_para_format(pf, after=0)
    pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(pf.add_run("可复用模板 | 按需定制"), size=9, color=MUTED)

    add_p(doc, "功能需求 PRD", before=8, after=2, size=13, bold=True, color=MUTED)
    p = add_p(doc, "项目管理系统", after=4, size=24, bold=True, color=INK)
    set_para_format(p, before=0, after=4, line=1.0)
    add_p(doc, "轻量、好用、刚需的项目执行与需求缺陷协同工具", after=14, size=13, color=MUTED)
    add_label_detail_table(
        doc,
        [
            ("文档类型", "功能需求 PRD"),
            ("适用对象", "产品、研发、测试、项目经理、团队负责人"),
            ("版本", "V1.0，可复用模板"),
            ("状态", "需求草案，可按业务场景裁剪"),
            ("核心模块", "仪表盘、甘特图、计划、需求、缺陷"),
        ],
    )
    rule = add_p(doc, "", after=12)
    paragraph_border_bottom(rule)


def build_doc():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_masthead(doc)

    add_callout(
        doc,
        "产品定位",
        "面向中小团队和项目制团队，提供项目概览、计划排期、需求跟踪、缺陷处理和进度提醒的一体化轻量项目管理系统。第一版聚焦刚需闭环，避免复杂审批、过度自定义和重型流程配置。",
    )

    add_p(doc, "1. 背景与目标", style="Heading 1")
    add_p(doc, "1.1 建设背景", style="Heading 2")
    add_p(
        doc,
        "团队在项目执行中常见信息分散、进度不可视、任务延期发现晚、需求与缺陷缺少关联、周报依赖人工整理等问题。系统需要用较低操作成本沉淀项目执行数据，让负责人随时看清整体情况，让执行人明确今日待办和风险事项。",
    )
    add_p(doc, "1.2 产品目标", style="Heading 2")
    add_bullets(
        doc,
        [
            "用一个仪表盘汇总项目健康度、执行进度、风险延期、需求缺陷和提交情况。",
            "用甘特图承载项目计划、里程碑、任务周期和延期影响，减少口头同步成本。",
            "用计划、需求、缺陷三类核心对象支撑项目执行闭环。",
            "提供提交模块和延期提醒，让执行进度可被持续更新、追踪和复盘。",
            "保留模板化能力，便于不同团队按项目类型复用和裁剪字段。",
        ],
    )
    add_p(doc, "1.3 设计原则", style="Heading 2")
    add_table(
        doc,
        ["原则", "说明", "落地要求"],
        [
            ["轻量", "减少流程负担，默认路径简单。", "新增事项不超过必要字段；列表支持快捷编辑；提醒规则默认可用。"],
            ["好用", "常用操作路径短，信息层级清晰。", "仪表盘先给结论；详情页再给字段；批量操作覆盖高频场景。"],
            ["刚需", "优先解决项目推进中的真实痛点。", "聚焦计划、需求、缺陷、延期、提交和统计，不做复杂 OA 审批。"],
            ["可复用", "支持团队沉淀项目模板。", "项目模板、字段模板、提醒模板、视图模板可复制复用。"],
        ],
        [1300, 3300, 4760],
    )

    add_p(doc, "2. 范围与角色", style="Heading 1")
    add_p(doc, "2.1 MVP 范围", style="Heading 2")
    add_table(
        doc,
        ["范围", "包含", "暂不包含"],
        [
            ["项目管理", "项目创建、成员、状态、周期、整体进度、风险统计。", "复杂项目组合管理、预算成本核算。"],
            ["计划管理", "计划拆解、负责人、起止时间、优先级、提交进度、延期提醒。", "复杂审批流、自动资源平衡。"],
            ["需求管理", "需求录入、优先级、状态流转、关联计划和缺陷。", "完整产品路线图和客户门户。"],
            ["缺陷管理", "缺陷提交、指派、状态、严重程度、关联需求和版本。", "自动化测试平台和代码扫描集成。"],
            ["可视化", "仪表盘、甘特图、基础报表、风险提醒。", "复杂 BI 自定义报表。"],
        ],
        [1600, 4200, 3560],
    )
    add_p(doc, "2.2 用户角色", style="Heading 2")
    add_table(
        doc,
        ["角色", "核心诉求", "关键权限"],
        [
            ["项目负责人", "看整体进度、风险和延期，推动成员更新。", "项目配置、成员管理、计划调整、提醒规则配置。"],
            ["产品/业务负责人", "管理需求优先级，跟踪交付范围。", "需求创建、需求排序、状态流转、需求验收。"],
            ["研发成员", "查看任务、提交进度、反馈风险。", "计划处理、进度提交、需求技术反馈、缺陷修复。"],
            ["测试成员", "提交缺陷、验证修复、查看版本范围。", "缺陷创建、指派、复测、关闭。"],
            ["团队管理者", "了解项目健康度和资源压力。", "查看汇总数据、导出报表、关注风险项目。"],
        ],
        [1700, 4200, 3460],
    )

    add_p(doc, "3. 信息架构与核心对象", style="Heading 1")
    add_p(doc, "3.1 功能模块", style="Heading 2")
    add_table(
        doc,
        ["模块", "定位", "核心能力"],
        [
            ["仪表盘", "项目驾驶舱。", "数据汇总、执行进度、风险延期、提交动态、待办提醒。"],
            ["甘特图", "时间和依赖视图。", "计划排期、里程碑、延期标识、任务拖拽调整、关键路径提示。"],
            ["计划", "项目执行拆解。", "阶段计划、任务、负责人、工期、提交模块、状态流转。"],
            ["需求", "交付范围管理。", "需求池、优先级、状态、关联计划、验收标准、版本归属。"],
            ["缺陷", "质量问题闭环。", "缺陷提交、指派修复、复测关闭、严重程度、关联需求。"],
        ],
        [1300, 2500, 5560],
    )
    add_p(doc, "3.2 核心数据对象", style="Heading 2")
    add_table(
        doc,
        ["对象", "关键字段", "关联关系"],
        [
            ["项目", "名称、状态、负责人、成员、开始日期、结束日期、健康度、完成率。", "包含计划、需求、缺陷、里程碑。"],
            ["计划/任务", "标题、类型、负责人、优先级、起止时间、进度、状态、延期原因。", "可关联需求、缺陷、里程碑。"],
            ["需求", "标题、来源、优先级、状态、验收标准、目标版本、业务价值。", "可拆解为计划，可关联缺陷。"],
            ["缺陷", "标题、严重程度、优先级、复现步骤、状态、修复人、验证人。", "可关联需求、计划、版本。"],
            ["提交记录", "提交人、提交对象、完成比例、说明、附件、风险、提交时间。", "绑定计划、需求或缺陷。"],
            ["提醒规则", "触发条件、提醒对象、提醒时间、提醒渠道、是否启用。", "作用于项目、计划、需求、缺陷。"],
        ],
        [1400, 4200, 3760],
    )

    add_p(doc, "4. 功能需求", style="Heading 1")
    add_p(doc, "4.1 仪表盘", style="Heading 2")
    add_p(doc, "仪表盘用于展示项目整体情况，默认作为进入系统后的首页。页面应先给出项目健康结论，再展示可追溯的数据来源。")
    add_table(
        doc,
        ["需求点", "功能说明", "优先级"],
        [
            ["项目概览卡", "展示项目总数、进行中项目、延期项目、已完成项目、风险项目。", "P0"],
            ["执行进度", "展示当前项目完成率、计划完成率、需求完成率、缺陷关闭率。", "P0"],
            ["风险提醒", "展示延期任务、临期任务、无人负责事项、长期未更新事项。", "P0"],
            ["提交动态", "展示成员最近提交的进度、风险说明、缺陷修复和需求验收记录。", "P0"],
            ["我的待办", "按用户聚合待处理计划、待确认需求、待修复缺陷、待复测缺陷。", "P0"],
            ["数据筛选", "支持按项目、负责人、时间范围、状态筛选仪表盘数据。", "P1"],
            ["报表导出", "支持导出当前项目概览为 Excel 或 PDF。", "P2"],
        ],
        [1700, 5860, 1800],
    )
    add_p(doc, "4.1.1 仪表盘指标口径", style="Heading 3")
    add_table(
        doc,
        ["指标", "计算口径", "展示建议"],
        [
            ["项目完成率", "已完成计划数量 / 全部计划数量，支持按权重扩展。", "百分比 + 趋势箭头。"],
            ["延期数量", "当前日期大于计划结束日期，且状态未完成的事项数量。", "红色数字 + 可点击列表。"],
            ["需求交付率", "已验收需求数量 / 已确认需求数量。", "条形进度。"],
            ["缺陷关闭率", "已关闭缺陷数量 / 全部有效缺陷数量。", "百分比 + 严重缺陷标识。"],
            ["提交活跃度", "近 7 天有提交记录的成员数量 / 项目成员数量。", "成员头像或数字。"],
        ],
        [1800, 4960, 2600],
    )

    add_p(doc, "4.2 甘特图", style="Heading 2")
    add_p(doc, "甘特图用于展示计划排期、任务跨度、里程碑和延期影响。第一版以清晰查看和轻量调整为主。")
    add_table(
        doc,
        ["需求点", "功能说明", "优先级"],
        [
            ["时间轴视图", "支持按日、周、月切换，默认按周展示。", "P0"],
            ["任务条展示", "展示计划名称、负责人、起止时间、状态、完成率。", "P0"],
            ["里程碑", "支持设置关键节点，里程碑在时间轴上突出显示。", "P0"],
            ["延期标识", "延期任务显示风险颜色，并展示延期天数。", "P0"],
            ["计划调整", "支持拖动任务条调整起止时间，变更需记录操作日志。", "P1"],
            ["依赖关系", "支持前置任务依赖，前置延期时提示受影响任务。", "P1"],
            ["视图过滤", "按负责人、状态、计划类型、是否延期过滤。", "P1"],
        ],
        [1700, 5860, 1800],
    )

    add_p(doc, "4.3 计划", style="Heading 2")
    add_p(doc, "计划模块是项目执行的主线，承担任务拆解、提交进度和延期管理。")
    add_table(
        doc,
        ["需求点", "功能说明", "优先级"],
        [
            ["计划列表", "按阶段、负责人、状态、优先级、起止时间展示计划。", "P0"],
            ["计划详情", "展示描述、目标、负责人、关联需求、关联缺陷、提交记录。", "P0"],
            ["提交模块", "成员可提交完成比例、进展说明、风险、附件和下一步计划。", "P0"],
            ["状态流转", "未开始、进行中、阻塞、待确认、已完成、已取消。", "P0"],
            ["延期提醒", "临期前提醒、到期未完成提醒、延期持续提醒。", "P0"],
            ["批量编辑", "支持批量修改负责人、状态、优先级和计划时间。", "P1"],
            ["计划模板", "支持从模板创建阶段计划和常用任务清单。", "P1"],
        ],
        [1700, 5860, 1800],
    )
    add_p(doc, "4.3.1 提交模块字段", style="Heading 3")
    add_table(
        doc,
        ["字段", "是否必填", "说明"],
        [
            ["完成比例", "是", "0% 到 100%，用于更新计划进度。"],
            ["进展说明", "是", "简要说明本次完成内容。"],
            ["风险/阻塞", "否", "存在延期风险或依赖阻塞时填写。"],
            ["下一步计划", "否", "说明下一次提交前计划完成内容。"],
            ["附件/链接", "否", "可上传文档、截图或外部链接。"],
        ],
        [2200, 1600, 5560],
    )

    add_p(doc, "4.4 需求", style="Heading 2")
    add_p(doc, "需求模块用于承载项目交付范围，强调需求状态、优先级、验收标准和计划关联。")
    add_table(
        doc,
        ["需求点", "功能说明", "优先级"],
        [
            ["需求池", "展示全部需求，支持按状态、优先级、版本、负责人筛选。", "P0"],
            ["需求详情", "包含背景、目标、验收标准、附件、评论和变更记录。", "P0"],
            ["状态流转", "草稿、待评审、已确认、开发中、待验收、已验收、已关闭。", "P0"],
            ["优先级", "支持 P0、P1、P2、P3，默认 P2。", "P0"],
            ["关联计划", "需求可关联一个或多个计划，用于追踪交付进度。", "P0"],
            ["需求变更", "关键字段变更需记录变更人、变更时间和变更原因。", "P1"],
            ["需求模板", "支持配置不同类型需求的字段和验收清单。", "P1"],
        ],
        [1700, 5860, 1800],
    )

    add_p(doc, "4.5 缺陷", style="Heading 2")
    add_p(doc, "缺陷模块用于质量问题闭环，强调快速提交、清晰指派、复测关闭和质量统计。")
    add_table(
        doc,
        ["需求点", "功能说明", "优先级"],
        [
            ["缺陷列表", "展示标题、严重程度、状态、负责人、创建人、关联需求。", "P0"],
            ["缺陷提交", "填写环境、复现步骤、实际结果、期望结果、截图或附件。", "P0"],
            ["状态流转", "新建、已确认、修复中、待复测、已关闭、已拒绝、重新打开。", "P0"],
            ["严重程度", "阻塞、严重、一般、轻微。", "P0"],
            ["指派与提醒", "缺陷创建后可指派修复人，到期未处理自动提醒。", "P0"],
            ["复测记录", "测试人员可提交复测结论，失败后重新打开。", "P0"],
            ["质量统计", "按严重程度、模块、负责人、版本统计缺陷数量和关闭率。", "P1"],
        ],
        [1700, 5860, 1800],
    )

    add_p(doc, "5. 通用能力", style="Heading 1")
    add_p(doc, "5.1 搜索与筛选", style="Heading 2")
    add_bullets(
        doc,
        [
            "全局搜索支持按标题、编号、负责人、标签检索计划、需求和缺陷。",
            "列表筛选支持保存为个人视图，例如“我的延期事项”“本周待完成”“P0 需求”。",
            "默认列表按优先级、延期状态和更新时间排序，保证重要事项靠前。",
        ],
    )
    add_p(doc, "5.2 提醒规则", style="Heading 2")
    add_table(
        doc,
        ["提醒类型", "触发条件", "提醒对象", "默认策略"],
        [
            ["临期提醒", "计划或缺陷距离截止时间小于 N 天。", "负责人、项目负责人。", "默认提前 1 天提醒。"],
            ["延期提醒", "超过截止时间且状态未完成。", "负责人、项目负责人。", "每日提醒一次，直到完成或改期。"],
            ["未更新提醒", "进行中事项超过 N 天无提交记录。", "负责人。", "默认 3 天提醒。"],
            ["阻塞提醒", "计划状态变为阻塞。", "项目负责人、关联成员。", "立即提醒。"],
            ["复测提醒", "缺陷进入待复测状态。", "验证人、创建人。", "立即提醒。"],
        ],
        [1600, 3500, 2200, 2060],
    )
    add_p(doc, "5.3 权限", style="Heading 2")
    add_table(
        doc,
        ["权限项", "项目负责人", "成员", "只读成员"],
        [
            ["查看项目", "允许", "允许", "允许"],
            ["配置项目", "允许", "不允许", "不允许"],
            ["创建计划", "允许", "允许，可配置", "不允许"],
            ["提交进度", "允许", "允许", "不允许"],
            ["管理需求", "允许", "允许，可配置", "不允许"],
            ["管理缺陷", "允许", "允许，可配置", "不允许"],
            ["导出数据", "允许", "允许，可配置", "不允许"],
        ],
        [2400, 2200, 2500, 2260],
    )

    add_p(doc, "6. 模板化与按需定制", style="Heading 1")
    add_p(doc, "模板能力用于复用项目管理经验，但第一版不做过度复杂的低代码配置。")
    add_table(
        doc,
        ["模板类型", "可配置内容", "典型用途"],
        [
            ["项目模板", "默认阶段、里程碑、计划清单、角色权限、提醒规则。", "新建研发项目、交付项目、专项项目。"],
            ["字段模板", "需求和缺陷的自定义字段、必填规则、默认值。", "适配不同团队的信息采集习惯。"],
            ["视图模板", "列表列配置、筛选条件、排序规则。", "保存“我的待办”“版本缺陷”等常用视图。"],
            ["提醒模板", "临期、延期、未更新、阻塞提醒的触发时间和对象。", "不同项目采用不同提醒强度。"],
            ["提交模板", "提交字段、提交频率、风险说明格式。", "统一周报、日报和阶段汇报口径。"],
        ],
        [1600, 4700, 3060],
    )
    add_p(doc, "6.1 推荐模板结构", style="Heading 2")
    add_numbered(
        doc,
        [
            "选择项目类型：研发项目、客户交付项目、运营专项、内部工具建设。",
            "生成默认模块：仪表盘、甘特图、计划、需求、缺陷。",
            "生成默认计划：启动、需求确认、设计、开发、测试、上线、复盘。",
            "生成默认提醒：临期提醒、延期提醒、未更新提醒、阻塞提醒。",
            "允许项目负责人在创建后删除不需要的字段和阶段。",
        ],
    )

    add_p(doc, "7. 非功能需求", style="Heading 1")
    add_table(
        doc,
        ["类别", "要求"],
        [
            ["易用性", "新增计划、需求、缺陷应在 1 分钟内完成；列表支持快捷筛选和批量操作。"],
            ["性能", "常规列表 1000 条以内筛选和分页响应不超过 1 秒；仪表盘首屏加载不超过 2 秒。"],
            ["安全", "按项目隔离数据；关键操作记录日志；导出能力可按角色控制。"],
            ["兼容", "支持主流桌面浏览器；移动端至少支持查看、提交进度和处理待办。"],
            ["可维护", "模块和字段配置应由后端配置驱动，避免每个项目写死字段。"],
        ],
        [1800, 7560],
    )

    add_p(doc, "8. 关键流程", style="Heading 1")
    add_p(doc, "8.1 新建项目流程", style="Heading 2")
    add_numbered(
        doc,
        [
            "项目负责人创建项目，填写名称、周期、负责人、成员和项目类型。",
            "选择项目模板，系统自动生成阶段计划、默认视图和提醒规则。",
            "负责人调整甘特图排期，设置里程碑和关键任务负责人。",
            "成员开始执行并按计划提交进度。",
            "仪表盘持续汇总进度、延期和提交情况。",
        ],
    )
    add_p(doc, "8.2 延期处理流程", style="Heading 2")
    add_numbered(
        doc,
        [
            "系统识别任务到期未完成，将任务标记为延期。",
            "系统向负责人和项目负责人发送延期提醒。",
            "负责人提交延期原因、当前进展和预计完成时间。",
            "项目负责人确认是否调整计划时间，变更记录进入日志。",
            "仪表盘和甘特图同步更新延期状态和影响范围。",
        ],
    )

    add_p(doc, "9. 验收标准", style="Heading 1")
    add_table(
        doc,
        ["模块", "验收标准"],
        [
            ["仪表盘", "可正确展示项目数量、完成率、延期数、需求交付率、缺陷关闭率和提交动态。"],
            ["甘特图", "可按时间展示计划、里程碑和延期任务，筛选和缩放正常。"],
            ["计划", "支持计划创建、编辑、状态流转、提交进度和延期提醒。"],
            ["需求", "支持需求创建、状态流转、优先级管理、关联计划和验收记录。"],
            ["缺陷", "支持缺陷提交、指派、修复、复测、关闭和重新打开。"],
            ["模板", "可基于模板创建项目，并允许删除或调整默认字段和阶段。"],
        ],
        [1800, 7560],
    )

    add_p(doc, "10. 后续迭代方向", style="Heading 1")
    add_bullets(
        doc,
        [
            "接入代码仓库、CI/CD 或测试平台，自动同步提交、构建和缺陷数据。",
            "增加项目周报自动生成能力，基于提交记录和延期数据生成汇总。",
            "支持跨项目资源视图，观察成员任务负载和延期集中点。",
            "支持更细粒度的工作流配置，但保持默认路径简单可用。",
        ],
    )

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build_doc()
    print(path.resolve())
